"""Local extraction with source-owned quotes; no model or filesystem access.

Limits bound common oversized inputs, not every adversarial parser workload.
Run behind request/time/memory limits before exposing uploads publicly.
"""
from io import BytesIO
from pathlib import PurePosixPath
import re
from typing import Literal
from zipfile import ZipFile
from xml.etree import ElementTree

from .models import Document, Evidence, Locator

SUPPORTED_EXTENSIONS = {".docx", ".xlsx", ".pdf", ".md", ".txt"}
MAX_INPUT_BYTES = 15 * 1024 * 1024
MAX_ZIP_BYTES = 40 * 1024 * 1024
MAX_ZIP_ENTRIES = 1500
MAX_PDF_PAGES = 200
MAX_CELLS = 100_000
MAX_EVIDENCE = 10_000
MAX_TEXT_CHARS = 2_000_000


class ExtractionLimit(Exception):
    pass


class _Collector:
    def __init__(self):
        self.items: list[tuple[str, Locator]] = []
        self.characters = 0
        self.warnings: list[str] = []

    def warn(self, message: str):
        if message not in self.warnings:
            self.warnings.append(message)

    def add(self, text: str, locator: Locator):
        if not text.strip():
            return
        if len(self.items) >= MAX_EVIDENCE or self.characters + len(text) > MAX_TEXT_CHARS:
            raise ExtractionLimit("Превышен лимит извлекаемого текста или фрагментов.")
        self.items.append((text, locator))
        self.characters += len(text)


def _inspect_zip(data: bytes) -> tuple[list[str], list[str]]:
    """Check archive metadata and actual streamed expansion before Office parsers."""
    warnings = []
    with ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ZIP_ENTRIES or sum(i.file_size for i in entries) > MAX_ZIP_BYTES:
            raise ExtractionLimit("Превышен лимит распакованного Office-документа.")
        total = 0
        names = []
        for entry in entries:
            name = entry.filename
            if name in names or entry.flag_bits & 1:
                raise ValueError("ambiguous or encrypted archive")
            names.append(name)
            if entry.file_size > 1024 * 1024 and entry.file_size > max(1, entry.compress_size) * 250:
                raise ExtractionLimit("Слишком высокая степень сжатия Office-документа.")
            with archive.open(entry) as stream:
                while chunk := stream.read(65536):
                    total += len(chunk)
                    if total > MAX_ZIP_BYTES:
                        raise ExtractionLimit("Превышен лимит распакованного Office-документа.")
            if name.endswith(".xml"):
                xml = archive.read(entry)
                if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                    raise ValueError("XML entities are not supported")
                if name == "word/document.xml":
                    if any(tag in xml for tag in (b"<w:ins ", b"<w:del ", b"<w:txbxContent", b"<w:altChunk")):
                        warnings.append("В Word есть исправления, текстовые поля или встроенные блоки; их содержимое может быть извлечено не полностью.")
                if name.startswith(("word/header", "word/footer", "word/footnotes", "word/endnotes")):
                    # PAGE fields and separator-only footnotes do not omit source prose.
                    tree = ElementTree.fromstring(xml)
                    if any((node.text or "").strip() for node in tree.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")):
                        warnings.append("Колонтитулы и сноски Word содержат текст, не включённый в извлечение; проверьте их отдельно.")
        if any("/media/" in name or "/drawings/" in name for name in names):
            warnings.append("В документе есть изображения или графические объекты. Их содержимое не распознано; OCR не выполнялся.")
    return names, warnings


def _docx(data: bytes, out: _Collector):
    from docx import Document as WordDocument
    from docx.text.paragraph import Paragraph

    _, warnings = _inspect_zip(data)
    for warning in warnings:
        out.warn(warning)
    document = WordDocument(BytesIO(data))
    section = None
    paragraph_index = 0
    table_index = 0
    cells_seen = 0
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph_index += 1
            if block.style and block.style.name.startswith("Heading"):
                section = block.text
            out.add(block.text, Locator(kind="paragraph", section=section, paragraph_index=paragraph_index))
        else:
            table_index += 1
            seen = set()
            for row_index, row in enumerate(block.rows, 1):
                for column_index, cell in enumerate(row.cells, 1 + row.grid_cols_before):
                    cells_seen += 1
                    if cells_seen > MAX_CELLS:
                        raise ExtractionLimit("Превышен лимит ячеек Word.")
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    if cell.tables:
                        out.warn("Вложенные таблицы Word не извлечены; проверьте их отдельно.")
                    location = f"Таблица {table_index}"
                    if section:
                        location = f"{section} / {location}"
                    out.add(cell.text, Locator(kind="table_cell", section=location, cell_range=f"R{row_index}C{column_index}"))


def _xlsx(data: bytes, out: _Collector):
    from openpyxl import load_workbook

    _, warnings = _inspect_zip(data)
    for warning in warnings:
        out.warn(warning)
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        count = 0
        for sheet in workbook.worksheets:
            if (sheet.max_row or 0) * (sheet.max_column or 0) + count > MAX_CELLS:
                raise ExtractionLimit("Превышен лимит ячеек Excel.")
            for row in sheet.iter_rows():
                for cell in row:
                    count += 1
                    if count > MAX_CELLS:
                        raise ExtractionLimit("Превышен лимит ячеек Excel.")
                    if cell.value is None:
                        continue
                    if cell.data_type == "f":
                        out.warn("В Excel есть формулы: сохранён текст формулы, вычисление и восстановление отображаемого значения не выполнялись.")
                    out.add(str(cell.value), Locator(kind="sheet_range", sheet=sheet.title, cell_range=cell.coordinate))
    finally:
        workbook.close()


def _pdf(data: bytes, out: _Collector):
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data), strict=False)
    if reader.is_encrypted:
        raise ValueError("encrypted PDF")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ExtractionLimit("Превышен лимит страниц PDF.")
    for index, page in enumerate(reader.pages, 1):
        try:
            text = page.extract_text() or ""
            if not text.strip():
                out.warn(f"Страница {index}: текст не извлечён; возможен скан или пустая страница. OCR не выполнялся.")
                continue
            resources = page.get("/Resources")
            resources = resources.get_object() if resources else {}
            objects = resources.get("/XObject")
            objects = objects.get_object() if objects else {}
            if any(obj.get_object().get("/Subtype") in ("/Image", "/Form") for obj in objects.values()):
                out.warn(f"Страница {index}: есть изображения или графические объекты; полнота их текстового содержания не проверена, OCR не выполнялся.")
            for fragment_index, fragment in enumerate(re.split(r"\n\s*\n", text), 1):
                out.add(fragment, Locator(kind="pdf_page", page=index, paragraph_index=fragment_index))
        except ExtractionLimit:
            raise
        except Exception:
            out.warn(f"Страница {index}: ошибка извлечения; содержимое страницы требует ручной проверки.")


def _text(data: bytes, out: _Collector):
    text = data.decode("utf-8-sig")
    for index, paragraph in enumerate(re.split(r"\r?\n[ \t]*\r?\n", text), 1):
        out.add(paragraph, Locator(kind="paragraph", paragraph_index=index))


def extract_document(
    data: bytes,
    filename: str,
    version: Literal["before", "after"],
    document_id: str,
) -> tuple[Document, list[Evidence]]:
    """Return evidence or an explicit safe failure; never treat failure as empty success."""
    name = PurePosixPath(filename.replace("\\", "/")).name
    extension = PurePosixPath(name).suffix.lower()
    out = _Collector()
    try:
        if extension not in SUPPORTED_EXTENSIONS:
            raise ExtractionLimit("Формат не поддерживается. Допустимы DOCX, XLSX, PDF, MD и TXT.")
        if not data or len(data) > MAX_INPUT_BYTES:
            raise ExtractionLimit("Документ пуст или превышает лимит размера 15 МиБ.")
        {".docx": _docx, ".xlsx": _xlsx, ".pdf": _pdf, ".md": _text, ".txt": _text}[extension](data, out)
        if not out.items:
            return Document(document_id=document_id, version=version, name=name, extraction_status="failed", warnings=out.warnings + ["Читаемый текст не найден. Документ не включён в анализ."]), []
    except ExtractionLimit as error:
        return Document(document_id=document_id, version=version, name=name, extraction_status="failed", warnings=[str(error)]), []
    except Exception:
        return Document(document_id=document_id, version=version, name=name, extraction_status="failed", warnings=["Не удалось прочитать документ: файл повреждён, зашифрован или его структура не поддерживается."]), []

    evidence = []
    for index, (quote, locator) in enumerate(out.items):
        neighbors = []
        if index:
            neighbors.append("Предыдущий фрагмент: " + out.items[index - 1][0][-700:])
        if index + 1 < len(out.items):
            neighbors.append("Следующий фрагмент: " + out.items[index + 1][0][:700])
        evidence.append(Evidence(
            evidence_id=f"{document_id}:e{index + 1:05d}", document_id=document_id,
            version=version, document_name=name, locator=locator, quote=quote,
            context="\n".join(neighbors),
            extraction_warning=("Документ извлечён частично. Проверьте предупреждения документа." if out.warnings else None),
        ))
    return Document(document_id=document_id, version=version, name=name,
                    extraction_status="partial" if out.warnings else "read", warnings=out.warnings), evidence
